import math
import re
from dataclasses import dataclass
from datetime import timedelta
from pathlib import Path
from typing import Optional

import openpyxl


@dataclass(frozen=True)
class GynSuggestion:
    name: str
    address: str
    city: Optional[str] = None
    phone: Optional[str] = None
    website: Optional[str] = None
    emails: Optional[str] = None
    rating: Optional[float] = None
    reviews: Optional[int] = None
    distance_km: Optional[float] = None


def suggest_top3(
    *,
    city: str | None = None,
    address_hint: str | None = None,
    latitude: float | None = None,
    longitude: float | None = None,
    xlsx_path: str | None = None,
) -> list[GynSuggestion]:
    """
    Suggests up to 3 gynecologists using city, address text and/or latitude-longitude.
    Works if the user provides at least one of: city, address_hint, latitude+longitude.
    """
    path = Path(xlsx_path or Path(__file__).resolve().parents[1] / "ginecologhe_donna_filtrate.xlsx")
    if not path.is_file():
        return []

    rows = _load_rows(path)
    if not rows:
        return []

    city_norm = _norm(city)
    address_norm = _norm(address_hint)
    has_coords = latitude is not None and longitude is not None

    if not city_norm and not address_norm and not has_coords:
        return []

    scored = []
    for item in rows:
        score = 0.0
        distance_km = None

        item_address_norm = _norm(item["address"])

        if city_norm and city_norm in item_address_norm:
            score += 100

        if address_norm:
            score += _token_overlap(address_norm, item_address_norm) * 80

        if has_coords and item["latitude"] is not None and item["longitude"] is not None:
            distance_km = _haversine_km(latitude, longitude, item["latitude"], item["longitude"])
            score += max(0, 150 - distance_km)

        rating = item["rating"] or 0
        reviews = item["reviews"] or 0
        score += rating * 5
        score += min(reviews, 200) / 20

        if score > 0:
            scored.append((score, distance_km, item))

    scored.sort(
        key=lambda x: (
            x[1] is None,
            x[1] if x[1] is not None else 999999,
            -x[0],
        )
    )

    return [
        GynSuggestion(
            name=item["name"],
            address=item["address"],
            city=item["city"],
            phone=item["phone"],
            website=item["website"],
            emails=item["emails"],
            rating=item["rating"],
            reviews=item["reviews"],
            distance_km=round(distance_km, 1) if distance_km is not None else None,
        )
        for _, distance_km, item in scored[:3]
    ]


def _load_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []

    headers = [
        str(cell).strip() if cell is not None else ""
        for cell in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    ]
    col = {name: idx for idx, name in enumerate(headers)}

    for row in ws.iter_rows(min_row=2, values_only=True):
        name = _str(row[col["Business Name"]])
        phone = _str(row[col["Number"]])
        address = _str(row[col["Address (Zip code, City, Country)"]])
        city = _str(row[col["Città"]])
        website = _str(row[col["Website"]])
        emails = _str(row[col["Emails"]])
        reviews = _float(row[col["Reviews"]])
        rating = _float(row[col["Rating"]])
        latitude = _coord(row[col["Latitude valida"]])
        longitude = _coord(row[col["Longitude valida"]])

        if not name or not address:
            continue

        out.append(
            {
                "name": name,
                "city": city,
                "phone": phone,
                "address": address,
                "website": website,
                "emails": emails,
                "reviews": int(reviews) if reviews is not None else None,
                "rating": rating,
                "latitude": latitude,
                "longitude": longitude,
            }
        )

    return out


def _str(value: object) -> Optional[str]:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _float(value: object) -> Optional[float]:
    if value is None:
        return None
    try:
        return float(str(value).strip())
    except Exception:
        return None


def _coord(value: object) -> Optional[float]:
    if value is None:
        return None
    if isinstance(value, timedelta):
        return value.days + value.seconds / 86400
    try:
        return float(value)
    except Exception:
        return None


def _norm(value: str | None) -> str:
    text = (value or "").lower()
    text = re.sub(r"[^a-z0-9àèéìòù]+", " ", text)
    return " ".join(text.split())


def _token_overlap(a: str, b: str) -> float:
    ta = {t for t in a.split() if len(t) >= 3}
    tb = {t for t in b.split() if len(t) >= 3}
    if not ta or not tb:
        return 0.0
    return len(ta & tb) / len(ta)


def _haversine_km(lat1: float, lon1: float, lat2: float, lon2: float) -> float:
    radius = 6371.0
    p1 = math.radians(lat1)
    p2 = math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)

    a = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return radius * 2 * math.atan2(math.sqrt(a), math.sqrt(1 - a))
